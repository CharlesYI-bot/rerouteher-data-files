from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "00_SOURCE_ARCHIVE"


def preview_csv(path: Path) -> None:
    df = pd.read_csv(path, nrows=5)
    print(f"\nCSV {path.name}: {list(df.columns)}")
    print(df.head(2).to_string(index=False))


def preview_xlsx(path: Path) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"\nXLSX {path.name}: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5), values_only=True))
        print(f"  SHEET {sheet_name!r}: rows={ws.max_row}, cols={ws.max_column}")
        for row in rows:
            print("   ", row)


for xlsx in [
    SRC / "ESCO_to_ONET-SOC_official.xlsx",
    SRC / "ESCO_v1.2.1_skills_occupations_matrix.xlsx",
]:
    preview_xlsx(xlsx)

for csv_path in sorted((SRC / "ONET_31.0").glob("*.csv")):
    preview_csv(csv_path)

for csv_path in [
    SRC / "Dingel_Neiman_occupations_workathome.csv",
    SRC / "GPTs_are_GPTs_occ_level.csv",
]:
    preview_csv(csv_path)
