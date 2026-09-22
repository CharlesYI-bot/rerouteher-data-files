#!/usr/bin/env python3
"""Normalize official NOC and OSCA downloads for the workbook builder."""

from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT = SCRIPT_DIR.parents[2]
SOURCE = PROJECT / "02_RAW_SOURCE_DATA" / "ITERATION_3" / "INTERVIEW_ASSISTANT"
NORMALIZED = PROJECT / "03_PROCESSED_DATA" / "ITERATION_3" / "INTERVIEW_ASSISTANT" / "normalized"


def tidy(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t\r\n;,.\u00a0")


def split_items(value: object) -> list[str]:
    text = tidy(value)
    if not text:
        return []
    parts = re.split(r"\s*;\s*|\n+|\s+•\s+", text)
    return [item for item in (tidy(part) for part in parts) if item]


def extract_osca() -> list[dict[str, object]]:
    workbook = load_workbook(SOURCE / "OSCA" / "osca_2024_descriptions.xlsx", read_only=True, data_only=True)
    sheet = workbook["Table 1"]
    roles: list[dict[str, object]] = []
    for row in sheet.iter_rows(min_row=6, values_only=True):
        code = tidy(row[0])
        if not re.fullmatch(r"\d{6}", code):
            continue
        tasks = split_items(row[8])
        specialisations = split_items(row[9])
        roles.append(
            {
                "code": code,
                "title": tidy(row[1]),
                "alternative_title": tidy(row[2]),
                "description": tidy(row[3]),
                "registration_or_licensing": tidy(row[4]),
                "inclusion_exclusion": tidy(row[5]),
                "skill_attributes": tidy(row[6]),
                "skill_level": tidy(row[7]),
                "tasks": tasks,
                "specialisations": specialisations,
            }
        )
    return roles


def extract_noc() -> list[dict[str, object]]:
    grouped: dict[str, dict[str, object]] = {}
    with (SOURCE / "NOC" / "noc_2021_v1_elements.csv").open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if row.get("Level") != "5":
                continue
            code = tidy(row.get("Code - NOC 2021 V1.0"))
            if not re.fullmatch(r"\d{5}", code):
                continue
            record = grouped.setdefault(
                code,
                {
                    "code": code,
                    "title": tidy(row.get("Class title")),
                    "main_duties": [],
                    "employment_requirements": [],
                    "illustrative_examples": [],
                    "inclusions": [],
                    "exclusions": [],
                    "additional_information": [],
                },
            )
            label = tidy(row.get("Element Type Label English"))
            value = tidy(row.get("Element Description English"))
            if not value:
                continue
            target = {
                "Main duties": "main_duties",
                "Employment requirements": "employment_requirements",
                "Illustrative example(s)": "illustrative_examples",
                "Inclusion(s)": "inclusions",
                "Exclusion(s)": "exclusions",
                "Additional information": "additional_information",
            }.get(label)
            if target:
                cast_list = record[target]
                assert isinstance(cast_list, list)
                cast_list.append(value)
    return list(grouped.values())


def main() -> None:
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    osca = extract_osca()
    noc = extract_noc()
    (NORMALIZED / "osca_2024_roles.json").write_text(
        json.dumps(osca, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (NORMALIZED / "noc_2021_v1_roles.json").write_text(
        json.dumps(noc, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"osca_roles": len(osca), "noc_roles": len(noc)}, indent=2))


if __name__ == "__main__":
    main()
